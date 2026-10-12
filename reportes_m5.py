"""Presentación Excel/PDF; los datos filtrados proceden del modelo Analisis."""

from datetime import timedelta
from html import escape
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer

CABECERAS = [
    "Ticket",
    "Asunto",
    "Categoría",
    "Área",
    "Prioridad",
    "Estado",
    "Solicitante",
    "Técnico",
    "Creado (Perú)",
    "SLA",
]


def texto_limpio(valor):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(valor))


def valores_ticket(ticket):
    fecha = (ticket["created_at"] - timedelta(hours=5)).strftime("%d/%m/%Y %H:%M")
    return [
        ticket["id"],
        ticket["asunto"],
        ticket["categoria"],
        ticket["area"],
        ticket["prioridad"],
        ticket["estado"],
        ticket["solicitante"],
        ticket["tecnico"] or "Sin asignar",
        fecha,
        ticket["sla"]["texto"],
    ]


def descripcion_filtros(filtros):
    return (
        " · ".join(f"{k}: {v}" for k, v in filtros.items() if v) or "Todos los tickets"
    )


def crear_excel(tickets, filtros):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Tickets"
    hoja.append(CABECERAS)
    for ticket in tickets:
        hoja.append(
            [
                texto_limpio(v) if isinstance(v, str) else v
                for v in valores_ticket(ticket)
            ]
        )
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, str):
                celda.value = texto_limpio(celda.value)
                celda.data_type = (
                    "s"  # Un asunto que inicia con '=' se exporta como texto.
                )
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="234BC4")
    for columna, ancho in zip("ABCDEFGHIJ", [10, 55, 20, 20, 14, 17, 26, 26, 23, 28]):
        hoja.column_dimensions[columna].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    criterios = libro.create_sheet("Criterios")
    criterios.append(["Filtros", descripcion_filtros(filtros)])
    criterios.append(["Total exportado", len(tickets)])
    criterios.append(
        ["Fechas", "Creación en hora de Perú (UTC-5), extremos inclusivos."]
    )
    criterios["B1"].data_type = "s"
    criterios.column_dimensions["A"].width = 22
    criterios.column_dimensions["B"].width = 95
    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


def crear_pdf(tickets, filtros):
    salida = BytesIO()
    estilos = getSampleStyleSheet()
    estilos.add(
        ParagraphStyle(
            name="CeldaReporte",
            fontName="Helvetica",
            fontSize=7,
            leading=10,
            wordWrap="CJK",
        )
    )

    def par(valor):
        return Paragraph(escape(texto_limpio(valor)), estilos["CeldaReporte"])

    filas = [[par(c) for c in CABECERAS]]
    filas.extend([par(v) for v in valores_ticket(t)] for t in tickets)
    if not tickets:
        filas.append([par("Sin tickets para los filtros seleccionados.")] + [""] * 9)
    tabla = Table(
        filas, colWidths=[32, 133, 60, 60, 40, 55, 75, 75, 78, 86], repeatRows=1
    )
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf0fc")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d5ddea")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    if not tickets:
        tabla.setStyle(TableStyle([("SPAN", (0, 1), (-1, 1))]))
    elementos = [
        Paragraph("Nexo TI · Reporte de tickets", estilos["Title"]),
        par(descripcion_filtros(filtros)),
        par(f"Total: {len(tickets)} · Fechas de creación en Perú (UTC-5)."),
        Spacer(1, 14),
        tabla,
    ]

    def pie(c, d):
        c.setFont("Helvetica", 8)
        c.drawString(24, 18, "Nexo TI · Reporte operativo")
        c.drawRightString(818, 18, f"Página {d.page}")

    SimpleDocTemplate(
        salida,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=32,
    ).build(elementos, onFirstPage=pie, onLaterPages=pie)
    salida.seek(0)
    return salida
