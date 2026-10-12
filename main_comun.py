"""Archivo común de M1–M5: configuración, funciones y rutas del primer avance.

M2 utiliza modelo y servicios en clases, controladores Flask y vistas Jinja.
Los demás módulos conservan sus funciones y SQL directo.
"""
import os
import re
import secrets
from datetime import datetime, timedelta, timezone
from functools import wraps
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pymysql.cursors
from PIL import Image, UnidentifiedImageError
from flask import (Flask, abort, flash, g, jsonify, redirect, render_template,
                   request, send_file, send_from_directory, session, url_for)
from werkzeug.security import check_password_hash, generate_password_hash


app = Flask(__name__)
Path(app.instance_path).mkdir(exist_ok=True)
archivo_clave = Path(app.instance_path) / 'clave_local.txt'
if not archivo_clave.exists():
    archivo_clave.write_text(secrets.token_hex(32), encoding='utf-8')
app.config.update(
    SECRET_KEY=os.environ.get('SECRET_KEY') or archivo_clave.read_text().strip(),
    DB_HOST=os.environ.get('DB_HOST', '127.0.0.1'),
    DB_PORT=int(os.environ.get('DB_PORT', '3306')),
    DB_USER=os.environ.get('DB_USER', 'root'),
    DB_PASSWORD=os.environ.get('DB_PASSWORD', ''),
    DB_NAME=os.environ.get('DB_NAME', 'soporte_ti'),
    UPLOAD_FOLDER=str(Path(app.root_path) / 'uploads'),
    MAX_CONTENT_LENGTH=16 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
)
ROLES = ('Administrador', 'Tecnico', 'Solicitante')
ESTADOS = ('Abierto', 'En proceso', 'Resuelto', 'Cerrado')
TRANSICIONES = {'Abierto': ('En proceso',), 'En proceso': ('Resuelto',),
               'Resuelto': ('Cerrado',), 'Cerrado': ()}


def ahora():
    # UTC en BD: no depende de la zona horaria del equipo ni del servidor MySQL.
    return datetime.now(timezone.utc).replace(tzinfo=None)


def obtener_conexion():
    return pymysql.connect(host=app.config['DB_HOST'], port=app.config['DB_PORT'],
                           user=app.config['DB_USER'], password=app.config['DB_PASSWORD'],
                           database=app.config['DB_NAME'], charset='utf8mb4',
                           cursorclass=pymysql.cursors.DictCursor, autocommit=False)


def consultar(sql, parametros=(), uno=False):
    with obtener_conexion() as objCon:
        with objCon.cursor() as cursor:
            cursor.execute(sql, parametros)
            return cursor.fetchone() if uno else cursor.fetchall()


def proteger(*roles):
    # Decorador pequeño: verificar sesión y rol antes de ejecutar la ruta.
    def decorador(funcion):
        @wraps(funcion)
        def ruta_protegida(*args, **kwargs):
            if not g.usuario:
                if request.path.startswith('/api/'):
                    abort(401)
                return redirect(url_for('login'))
            if roles and g.usuario['rol'] not in roles:
                abort(403)
            return funcion(*args, **kwargs)
        return ruta_protegida
    return decorador


@app.before_request
def cargar_usuario_y_validar_formulario():
    g.usuario = None
    if request.endpoint == 'static':
        return
    if request.path.startswith('/api/') and request.method not in ('GET', 'HEAD', 'OPTIONS'):
        abort(405)
    if session.get('usuario_id'):
        g.usuario = consultar('SELECT id, nombre, apellido, email, rol, activo FROM users WHERE id=%s',
                              (session['usuario_id'],), uno=True)
        if not g.usuario or not g.usuario['activo']:
            session.clear()
            g.usuario = None
    if 'csrf' not in session:
        session['csrf'] = secrets.token_hex(32)
    if request.method == 'POST':
        token = request.form.get('csrf_token', '')
        if not secrets.compare_digest(token.encode('utf-8'), session['csrf'].encode('utf-8')):
            abort(400, description='El formulario venció. Recarga la página e inténtalo nuevamente.')


@app.context_processor
def datos_comunes():
    cantidad = 0
    if getattr(g, 'usuario', None):
        cantidad = consultar('SELECT COUNT(*) AS cantidad FROM notificaciones WHERE usuario_id=%s AND leido=0',
                             (g.usuario['id'],), uno=True)['cantidad']
    return dict(usuario_actual=getattr(g, 'usuario', None), no_leidas=cantidad,
                csrf_token=session.get('csrf', ''), roles=ROLES, estados=ESTADOS)


@app.template_filter('fecha')
def fecha_peru(valor):
    return (valor - timedelta(hours=5)).strftime('%d/%m/%Y %H:%M') if valor else '—'


@app.template_filter('relativa')
def fecha_relativa(valor):
    minutos = max(0, int((ahora() - valor).total_seconds() / 60))
    if minutos < 1:
        return 'Hace un momento'
    if minutos < 60:
        return f'Hace {minutos} min'
    if minutos < 1440:
        return f'Hace {minutos // 60} h'
    return f'Hace {minutos // 1440} días'


def calcular_sla(ticket, momento=None):
    fin = ticket['fecha_resolucion'] or momento or ahora()
    total = (ticket['fecha_limite'] - ticket['created_at']).total_seconds()
    transcurrido = max(0, (fin - ticket['created_at']).total_seconds())
    porcentaje = transcurrido / total if total > 0 else 1
    if ticket['fecha_resolucion']:
        color = 'verde' if fin <= ticket['fecha_limite'] else 'rojo'
        texto = 'Resuelto a tiempo' if color == 'verde' else 'Resuelto fuera de plazo'
    elif porcentaje >= 1:
        color, texto = 'rojo', 'Vencido'
    elif porcentaje >= 0.75:
        color, texto = 'amarillo', 'Por vencer'
    else:
        color, texto = 'verde', 'En plazo'
    minutos = max(0, int((ticket['fecha_limite'] - fin).total_seconds() / 60))
    return dict(color=color, texto=texto, porcentaje=min(100, round(porcentaje * 100)),
                restante=f'{minutos // 60} h {minutos % 60} min')


def puede_ver(ticket):
    return ServicioTickets.puede_ver(ticket, g.usuario)


def obtener_ticket(ticket_id, cursor=None, bloquear=False):
    # Mantiene la interfaz que utilizan las rutas de M3.
    return acceso_m2(servicio_tickets().obtener, ticket_id, g.usuario, cursor, bloquear)


def notificar(cursor, ticket, tipo, mensaje, destinatarios=None):
    if destinatarios is None:
        destinatarios = [ticket['solicitante_id'], ticket['tecnico_id']]
    for usuario_id in set(destinatarios):
        if usuario_id and usuario_id != g.usuario['id']:
            cursor.execute('''INSERT INTO notificaciones
                (usuario_id, referencia_ticket, tipo, mensaje, created_at) VALUES (%s,%s,%s,%s,%s)''',
                           (usuario_id, ticket['id'], tipo, mensaje, ahora()))


def registrar_historial(cursor, ticket_id, tipo, comentario, anterior=None, nuevo=None,
                        tecnico_anterior=None, tecnico_nuevo=None):
    cursor.execute('''INSERT INTO ticket_historial (ticket_id, usuario_id, tipo, comentario,
        estado_anterior, estado_nuevo, tecnico_anterior_id, tecnico_nuevo_id, created_at)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
                   (ticket_id, g.usuario['id'], tipo, comentario, anterior, nuevo,
                    tecnico_anterior, tecnico_nuevo, ahora()))


@app.route('/')
def index():
    return redirect(url_for('tickets' if g.usuario else 'login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if g.usuario:
        return redirect(url_for('tickets'))
    errores = []
    if request.method == 'POST':
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        usuario = consultar('SELECT * FROM users WHERE email=%s', (email,), uno=True)
        if usuario and usuario['activo'] and check_password_hash(usuario['password_hash'], password):
            session.clear()
            session['usuario_id'] = usuario['id']
            session['csrf'] = secrets.token_hex(32)
            session.permanent = True
            return redirect(url_for('tickets'))
        errores.append('Correo o contraseña incorrectos, o cuenta inactiva.')
    return render_template('login_m1.html', errores=errores)


@app.route('/logout', methods=['POST'])
@proteger()
def logout():
    session.clear()
    return redirect(url_for('login'))


@app.route('/usuarios')
@proteger('Administrador')
def usuarios():
    lista = consultar('SELECT id, nombre, apellido, email, rol, activo FROM users ORDER BY id DESC')
    return render_template('usuarios_m1.html', usuarios=lista)


@app.route('/usuarios/nuevo', methods=['GET', 'POST'])
@app.route('/usuarios/<int:usuario_id>/editar', methods=['GET', 'POST'])
@proteger('Administrador')
def usuario_form(usuario_id=None):
    usuario = consultar('SELECT id, nombre, apellido, email, rol, activo FROM users WHERE id=%s',
                        (usuario_id,), uno=True) if usuario_id else None
    if usuario_id and not usuario:
        abort(404)
    errores = []
    if request.method == 'POST':
        nombre = request.form.get('nombre', '').strip()
        apellido = request.form.get('apellido', '').strip()
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        rol = request.form.get('rol', '')
        activo = 1 if request.form.get('activo') == '1' else 0
        if not 2 <= len(nombre) <= 50 or not 2 <= len(apellido) <= 50:
            errores.append('Nombre y apellido deben tener entre 2 y 50 caracteres.')
        if len(email) > 120 or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email):
            errores.append('Ingresa un correo válido de hasta 120 caracteres.')
        if (not usuario_id or password) and not 8 <= len(password) <= 128:
            errores.append('La contraseña debe tener entre 8 y 128 caracteres.')
        if password != request.form.get('confirm_password', ''):
            errores.append('Las contraseñas no coinciden.')
        if rol not in ROLES:
            errores.append('Selecciona un rol válido.')
        if usuario_id == g.usuario['id'] and (rol != 'Administrador' or not activo):
            errores.append('No puedes quitarte el acceso de administrador.')
        if not errores:
            try:
                with obtener_conexion() as objCon:
                    with objCon.cursor() as cursor:
                        if usuario_id:
                            cursor.execute('SELECT id FROM users WHERE id=%s FOR UPDATE', (usuario_id,))
                            cursor.execute("SELECT id FROM tickets WHERE tecnico_id=%s AND estado<>'Cerrado'", (usuario_id,))
                            if cursor.fetchone() and (rol != 'Tecnico' or not activo):
                                errores.append('Reasigna sus tickets pendientes y cierra los resueltos antes de desactivar o cambiar su rol.')
                            else:
                                cursor.execute('UPDATE users SET nombre=%s, apellido=%s, email=%s, rol=%s, activo=%s WHERE id=%s',
                                               (nombre, apellido, email, rol, activo, usuario_id))
                                if password:
                                    cursor.execute('UPDATE users SET password_hash=%s WHERE id=%s',
                                                   (generate_password_hash(password), usuario_id))
                        else:
                            cursor.execute('''INSERT INTO users (nombre,apellido,email,password_hash,rol,activo,created_at)
                                VALUES (%s,%s,%s,%s,%s,%s,%s)''',
                                           (nombre, apellido, email, generate_password_hash(password), rol, activo, ahora()))
                    if not errores:
                        objCon.commit()
                if not errores:
                    flash('Usuario guardado correctamente.', 'exito')
                    return redirect(url_for('usuarios'))
            except pymysql.err.IntegrityError:
                errores.append('Ya existe un usuario con ese correo.')
    return render_template('usuario_form_m1.html', usuario=usuario, errores=errores)


class ServicioFotos:
    """Validación y almacenamiento de evidencias, independiente de HTTP."""

    @staticmethod
    def preparar(archivos):
        fotos = []
        archivos = [a for a in archivos if a.filename]
        if len(archivos) > 3:
            raise ValueError('Puedes adjuntar un máximo de 3 fotos.')
        formatos = {'JPEG': ('image/jpeg', 'jpg'), 'PNG': ('image/png', 'png'), 'WEBP': ('image/webp', 'webp')}
        for archivo in archivos:
            datos = archivo.read(5 * 1024 * 1024 + 1)
            if len(datos) > 5 * 1024 * 1024:
                raise ValueError('Cada foto puede pesar hasta 5 MB.')
            try:
                with Image.open(BytesIO(datos)) as foto:
                    formato = foto.format
                    if formato not in formatos or archivo.mimetype != formatos[formato][0]:
                        raise ValueError('Solo se admiten imágenes JPG, PNG o WebP con tipo MIME correcto.')
                    if foto.width * foto.height > 20_000_000:
                        raise ValueError('La imagen supera el máximo de 20 megapíxeles.')
                    foto.verify()
                with Image.open(BytesIO(datos)) as foto:
                    foto.load()
                    copia = foto.convert('RGB') if formato == 'JPEG' else foto.convert('RGBA')
                    salida, mini = BytesIO(), BytesIO()
                    copia.save(salida, format=formato)
                    copia.thumbnail((160, 120))
                    copia.save(mini, format=formato)
                nombre = uuid4().hex + '.' + formatos[formato][1]
                fotos.append(dict(archivo=nombre, nombre_original=Path(archivo.filename).name[:255],
                                  mime=formatos[formato][0], datos=salida.getvalue(), mini=mini.getvalue()))
            except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
                raise ValueError('Uno de los archivos no es una imagen válida.') from exc
        return fotos

    def __init__(self, carpeta):
        self.carpeta = Path(carpeta)

    def guardar(self, fotos, guardados):
        self.carpeta.mkdir(exist_ok=True, parents=True)
        for foto in fotos:
            for prefijo, datos in [('', foto['datos']), ('mini_', foto['mini'])]:
                destino = self.carpeta / (prefijo + foto['archivo'])
                guardados.append(destino)
                destino.write_bytes(datos)

    @staticmethod
    def limpiar(guardados):
        for ruta in guardados:
            ruta.unlink(missing_ok=True)


class ModeloTickets:
    """Modelo: SQL y persistencia. No accede a request, session ni g."""

    @staticmethod
    def catalogos():
        return (consultar('SELECT * FROM categorias ORDER BY id'),
                consultar('SELECT * FROM prioridades ORDER BY id'))

    @staticmethod
    def obtener(ticket_id, cursor=None, bloquear=False):
        sql = 'SELECT * FROM tickets WHERE id=%s' + (' FOR UPDATE' if bloquear else '')
        if cursor:
            cursor.execute(sql, (ticket_id,))
            return cursor.fetchone()
        return consultar(sql, (ticket_id,), uno=True)

    @staticmethod
    def listar(usuario, filtros):
        condiciones, parametros = [], []
        if usuario['rol'] == 'Solicitante':
            condiciones.append('t.solicitante_id=%s')
            parametros.append(usuario['id'])
        elif usuario['rol'] == 'Tecnico':
            condiciones.append('(t.tecnico_id=%s OR t.solicitante_id=%s)')
            parametros.extend([usuario['id'], usuario['id']])
        estado = filtros.get('estado', '')
        prioridad = filtros.get('prioridad', '')
        busqueda = filtros.get('q', '').strip()[:150]
        if estado in ESTADOS:
            condiciones.append('t.estado=%s')
            parametros.append(estado)
        if prioridad in ('Alta', 'Media', 'Baja'):
            condiciones.append('p.nombre=%s')
            parametros.append(prioridad)
        if busqueda:
            condiciones.append('(t.asunto LIKE %s OR CAST(t.id AS CHAR)=%s)')
            parametros.extend(['%' + busqueda + '%', busqueda])
        sql = '''SELECT t.*, p.nombre AS prioridad, c.nombre AS categoria,
            CONCAT(u.nombre,' ',u.apellido) AS solicitante,
            CONCAT(te.nombre,' ',te.apellido) AS tecnico,
            (SELECT MIN(f.id) FROM ticket_fotos f WHERE f.ticket_id=t.id) AS foto_id,
            TIMESTAMPDIFF(SECOND, %s, t.fecha_limite) AS segundos_restantes
            FROM tickets t JOIN prioridades p ON p.id=t.prioridad_id
            JOIN categorias c ON c.id=t.categoria_id JOIN users u ON u.id=t.solicitante_id
            LEFT JOIN users te ON te.id=t.tecnico_id'''
        if condiciones:
            sql += ' WHERE ' + ' AND '.join(condiciones)
        sql += ' ORDER BY t.created_at DESC, t.id DESC'
        lista = consultar(sql, [ahora()] + parametros)
        return lista

    @staticmethod
    def detalle(ticket):
        resultado = dict(ticket)
        resultado['categoria'] = consultar('SELECT nombre FROM categorias WHERE id=%s', (ticket['categoria_id'],), True)['nombre']
        resultado['prioridad'] = consultar('SELECT nombre FROM prioridades WHERE id=%s', (ticket['prioridad_id'],), True)['nombre']
        resultado['solicitante'] = consultar("SELECT CONCAT(nombre,' ',apellido) AS nombre FROM users WHERE id=%s", (ticket['solicitante_id'],), True)['nombre']
        tecnico = consultar("SELECT CONCAT(nombre,' ',apellido) AS nombre FROM users WHERE id=%s", (ticket['tecnico_id'],), True)
        resultado['tecnico'] = tecnico['nombre'] if tecnico else None
        return resultado

    @staticmethod
    def fotos(ticket_id):
        return consultar('SELECT * FROM ticket_fotos WHERE ticket_id=%s ORDER BY id', (ticket_id,))

    @staticmethod
    def foto(foto_id):
        return consultar('SELECT * FROM ticket_fotos WHERE id=%s', (foto_id,), uno=True)

    @staticmethod
    def historial(ticket_id):
        return consultar('''SELECT h.*, CONCAT(u.nombre,' ',u.apellido) AS autor FROM ticket_historial h
            JOIN users u ON u.id=h.usuario_id WHERE h.ticket_id=%s ORDER BY h.created_at,h.id''', (ticket_id,))

    @staticmethod
    def tecnicos():
        return consultar("SELECT id,nombre,apellido FROM users WHERE rol='Tecnico' AND activo=1 ORDER BY nombre")

    @staticmethod
    def insertar(cursor, datos, solicitante_id, fecha, horas_sla):
        cursor.execute('''INSERT INTO tickets (asunto,descripcion,categoria_id,prioridad_id,
            solicitante_id,created_at,updated_at,fecha_limite) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)''',
                       (datos['asunto'], datos['descripcion'], datos['categoria_id'], datos['prioridad_id'],
                        solicitante_id, fecha, fecha, fecha + timedelta(hours=horas_sla)))
        return cursor.lastrowid

    @staticmethod
    def registrar_fotos(cursor, ticket_id, fotos, fecha):
        for foto in fotos:
            cursor.execute('''INSERT INTO ticket_fotos
                (ticket_id,nombre_original,archivo,mime,created_at) VALUES (%s,%s,%s,%s,%s)''',
                           (ticket_id, foto['nombre_original'], foto['archivo'], foto['mime'], fecha))

    @staticmethod
    def registrar_creacion(cursor, ticket_id, usuario_id, fecha):
        cursor.execute('''INSERT INTO ticket_historial
            (ticket_id,usuario_id,tipo,comentario,estado_nuevo,created_at)
            VALUES (%s,%s,'Registro','Ticket registrado.','Abierto',%s)''',
                       (ticket_id, usuario_id, fecha))


class ValidacionTicket(ValueError):
    def __init__(self, errores):
        self.errores = errores
        super().__init__(' '.join(errores))


class ServicioTickets:
    """Reglas de M2 y coordinación de BD/archivos, sin respuestas HTTP."""

    def __init__(self, modelo, fotos, conexion=obtener_conexion):
        self.modelo = modelo
        self.fotos = fotos
        self.conexion = conexion

    @staticmethod
    def puede_ver(ticket, usuario):
        return usuario['rol'] == 'Administrador' or ticket['solicitante_id'] == usuario['id'] or (
            usuario['rol'] == 'Tecnico' and ticket['tecnico_id'] == usuario['id'])

    def obtener(self, ticket_id, usuario, cursor=None, bloquear=False):
        ticket = self.modelo.obtener(ticket_id, cursor, bloquear)
        if not ticket:
            raise LookupError('Ticket inexistente.')
        if not self.puede_ver(ticket, usuario):
            raise PermissionError('Sin acceso al ticket.')
        return ticket

    def listar(self, usuario, filtros):
        lista = self.modelo.listar(usuario, filtros)
        for ticket in lista:
            ticket['sla'] = calcular_sla(ticket)
        return lista

    def crear(self, formulario, archivos, usuario, categorias, prioridades):
        datos = {k: formulario.get(k, '').strip() for k in ('asunto', 'descripcion')}
        for campo in ('categoria_id', 'prioridad_id'):
            try:
                datos[campo] = int(formulario.get(campo, ''))
            except (TypeError, ValueError):
                datos[campo] = None
        prioridad = next((p for p in prioridades if p['id'] == datos['prioridad_id']), None)
        errores = []
        if not 5 <= len(datos['asunto']) <= 150:
            errores.append('El asunto debe tener entre 5 y 150 caracteres.')
        if not 10 <= len(datos['descripcion']) <= 5000:
            errores.append('La descripción debe tener entre 10 y 5000 caracteres.')
        if datos['categoria_id'] not in [c['id'] for c in categorias] or not prioridad:
            errores.append('Selecciona categoría y prioridad válidas.')
        try:
            fotos = self.fotos.preparar(archivos)
        except ValueError as exc:
            errores.append(str(exc))
        if errores:
            raise ValidacionTicket(errores)
        guardados = []
        try:
            with self.conexion() as conexion:
                try:
                    with conexion.cursor() as cursor:
                        fecha = ahora()
                        ticket_id = self.modelo.insertar(cursor, datos, usuario['id'], fecha, prioridad['horas_sla'])
                        self.fotos.guardar(fotos, guardados)
                        self.modelo.registrar_fotos(cursor, ticket_id, fotos, fecha)
                        self.modelo.registrar_creacion(cursor, ticket_id, usuario['id'], fecha)
                    conexion.commit()
                except Exception:
                    conexion.rollback()
                    raise
        except Exception:
            self.fotos.limpiar(guardados)
            raise
        return ticket_id

    def detalle(self, ticket_id, usuario):
        ticket = self.modelo.detalle(self.obtener(ticket_id, usuario))
        ticket['sla'] = calcular_sla(ticket)
        return dict(ticket=ticket, fotos=self.modelo.fotos(ticket_id),
                    historial=self.modelo.historial(ticket_id),
                    tecnicos=self.modelo.tecnicos() if usuario['rol'] == 'Administrador' else [],
                    siguientes=TRANSICIONES[ticket['estado']],
                    puede_cambiar=usuario['rol'] == 'Tecnico' and ticket['tecnico_id'] == usuario['id'])

    def foto(self, foto_id, usuario):
        foto = self.modelo.foto(foto_id)
        if not foto:
            raise LookupError('Foto inexistente.')
        self.obtener(foto['ticket_id'], usuario)
        return foto


def servicio_tickets():
    # Composición de dependencias; respeta la configuración de cada entorno.
    return ServicioTickets(ModeloTickets(), ServicioFotos(app.config['UPLOAD_FOLDER']))


def acceso_m2(operacion, *args, **kwargs):
    # Traducción de errores de negocio a HTTP, responsabilidad del controlador.
    try:
        return operacion(*args, **kwargs)
    except LookupError:
        abort(404)
    except PermissionError:
        abort(403)


@app.route('/tickets')
@proteger()
def tickets():
    lista = servicio_tickets().listar(g.usuario, request.args)
    return render_template('tickets_m2_m3_m4.html', tickets=lista)


def preparar_fotos(archivos):
    # Compatibilidad con llamadas anteriores; la implementación pertenece al servicio.
    return ServicioFotos.preparar(archivos)


@app.route('/tickets/nuevo', methods=['GET', 'POST'])
@proteger()
def ticket_nuevo():
    servicio = servicio_tickets()
    categorias, prioridades = servicio.modelo.catalogos()
    errores = []
    if request.method == 'POST':
        try:
            ticket_id = servicio.crear(request.form, request.files.getlist('fotos'), g.usuario,
                                      categorias, prioridades)
        except ValidacionTicket as exc:
            errores = exc.errores
        else:
            flash('Ticket registrado. El área de TI podrá asignarlo a un técnico.', 'exito')
            return redirect(url_for('ticket_detalle', ticket_id=ticket_id))
    return render_template('ticket_form_m2.html', categorias=categorias, prioridades=prioridades, errores=errores)


@app.route('/tickets/<int:ticket_id>')
@proteger()
def ticket_detalle(ticket_id):
    datos = acceso_m2(servicio_tickets().detalle, ticket_id, g.usuario)
    return render_template('ticket_detalle_m2_m3_m4.html', **datos)


@app.route('/tickets/<int:ticket_id>/asignar', methods=['POST'])
@proteger('Administrador')
def asignar_ticket(ticket_id):
    tecnico_id = request.form.get('tecnico_id', type=int)
    comentario = request.form.get('comentario', '').strip()
    if not 3 <= len(comentario) <= 2000:
        abort(400, description='Describe el motivo de la asignación (3 a 2000 caracteres).')
    with obtener_conexion() as objCon:
        with objCon.cursor() as cursor:
            # Bloquear usuario antes de ticket, igual que en la edición de usuarios.
            cursor.execute("SELECT id FROM users WHERE id=%s AND rol='Tecnico' AND activo=1 FOR UPDATE", (tecnico_id,))
            if not cursor.fetchone():
                abort(400, description='Selecciona un técnico activo.')
            ticket = obtener_ticket(ticket_id, cursor, True)
            if ticket['estado'] not in ('Abierto', 'En proceso') or ticket['tecnico_id'] == tecnico_id:
                abort(400, description='El ticket ya tiene ese técnico o no admite asignación.')
            cursor.execute('UPDATE tickets SET tecnico_id=%s,updated_at=%s WHERE id=%s', (tecnico_id, ahora(), ticket_id))
            registrar_historial(cursor, ticket_id, 'Asignacion', comentario,
                                tecnico_anterior=ticket['tecnico_id'], tecnico_nuevo=tecnico_id)
            ticket['tecnico_id'] = tecnico_id
            notificar(cursor, ticket, 'Asignacion', f'Ticket #{ticket_id}: se asignó un técnico responsable.')
        objCon.commit()
    flash('Asignación guardada y notificaciones enviadas.', 'exito')
    return redirect(url_for('ticket_detalle', ticket_id=ticket_id))


@app.route('/tickets/<int:ticket_id>/estado', methods=['POST'])
@proteger('Tecnico')
def cambiar_estado(ticket_id):
    nuevo = request.form.get('estado', '')
    comentario = request.form.get('comentario', '').strip()
    with obtener_conexion() as objCon:
        with objCon.cursor() as cursor:
            ticket = obtener_ticket(ticket_id, cursor, True)
            if ticket['tecnico_id'] != g.usuario['id']:
                abort(403)
            if nuevo not in TRANSICIONES[ticket['estado']]:
                abort(400, description='Transición inválida. Actualiza el detalle del ticket.')
            if not 3 <= len(comentario) <= 2000:
                abort(400, description='El cambio de estado requiere un comentario de 3 a 2000 caracteres.')
            resolucion = ahora() if nuevo == 'Resuelto' else ticket['fecha_resolucion']
            cursor.execute('UPDATE tickets SET estado=%s,fecha_resolucion=%s,updated_at=%s WHERE id=%s',
                           (nuevo, resolucion, ahora(), ticket_id))
            registrar_historial(cursor, ticket_id, 'Estado', comentario, ticket['estado'], nuevo)
            notificar(cursor, ticket, 'Estado', f'Ticket #{ticket_id}: cambió a {nuevo}.')
        objCon.commit()
    flash('Estado actualizado.', 'exito')
    return redirect(url_for('ticket_detalle', ticket_id=ticket_id))


@app.route('/tickets/<int:ticket_id>/comentarios', methods=['POST'])
@proteger()
def agregar_comentario(ticket_id):
    comentario = request.form.get('comentario', '').strip()
    if not 3 <= len(comentario) <= 2000:
        abort(400, description='Escribe un comentario de 3 a 2000 caracteres.')
    with obtener_conexion() as objCon:
        with objCon.cursor() as cursor:
            ticket = obtener_ticket(ticket_id, cursor, True)
            if ticket['estado'] == 'Cerrado':
                abort(400, description='No se pueden agregar comentarios a un ticket cerrado.')
            registrar_historial(cursor, ticket_id, 'Comentario', comentario)
            notificar(cursor, ticket, 'Comentario', f'Ticket #{ticket_id}: recibiste un comentario.')
        objCon.commit()
    flash('Comentario publicado.', 'exito')
    return redirect(url_for('ticket_detalle', ticket_id=ticket_id))


@app.route('/fotos/<int:foto_id>')
@proteger()
def ver_foto(foto_id):
    foto = acceso_m2(servicio_tickets().foto, foto_id, g.usuario)
    nombre = ('mini_' if request.args.get('mini') == '1' else '') + foto['archivo']
    respuesta = send_from_directory(app.config['UPLOAD_FOLDER'], nombre, mimetype=foto['mime'])
    respuesta.headers['Cache-Control'] = 'private, no-store'
    respuesta.headers['X-Content-Type-Options'] = 'nosniff'
    return respuesta


@app.route('/sla')
@proteger()
def sla():
    return render_template('sla_m4.html', prioridades=consultar('SELECT * FROM prioridades ORDER BY id'))


@app.route('/notificaciones')
@proteger()
def notificaciones():
    lista = consultar('SELECT * FROM notificaciones WHERE usuario_id=%s ORDER BY created_at DESC,id DESC', (g.usuario['id'],))
    return render_template('notificaciones_m5.html', notificaciones=lista)


@app.route('/notificaciones/contador')
@proteger()
def contador_notificaciones():
    # M5: contador de notificaciones; la API de consulta también pertenece a M5.
    fila = consultar('SELECT COUNT(*) AS cantidad FROM notificaciones WHERE usuario_id=%s AND leido=0', (g.usuario['id'],), True)
    respuesta = jsonify(fila)
    respuesta.headers['Cache-Control'] = 'no-store'
    return respuesta


@app.route('/notificaciones/<int:notificacion_id>/leer', methods=['POST'])
@proteger()
def leer_notificacion(notificacion_id):
    with obtener_conexion() as objCon:
        with objCon.cursor() as cursor:
            cursor.execute('SELECT id FROM notificaciones WHERE id=%s AND usuario_id=%s', (notificacion_id, g.usuario['id']))
            if not cursor.fetchone():
                abort(404)
            cursor.execute('UPDATE notificaciones SET leido=1 WHERE id=%s AND usuario_id=%s', (notificacion_id, g.usuario['id']))
        objCon.commit()
    return redirect(url_for('notificaciones'))


@app.route('/notificaciones/leer-todas', methods=['POST'])
@proteger()
def leer_todas():
    with obtener_conexion() as objCon:
        with objCon.cursor() as cursor:
            cursor.execute('UPDATE notificaciones SET leido=1 WHERE usuario_id=%s AND leido=0', (g.usuario['id'],))
        objCon.commit()
    flash('Tus notificaciones se marcaron como leídas.', 'exito')
    return redirect(url_for('notificaciones'))


SQL_ANALISIS = ''' FROM tickets t
    JOIN prioridades p ON p.id=t.prioridad_id
    JOIN categorias c ON c.id=t.categoria_id
    JOIN users u ON u.id=t.solicitante_id
    LEFT JOIN users te ON te.id=t.tecnico_id'''


def filtros_analisis():
    # Compartido M4/M5: filtros de dashboard, reportes y API de consulta.
    filtros = {campo: request.args.get(campo, '').strip()
               for campo in ('desde', 'hasta', 'estado', 'prioridad')}
    condiciones, parametros, fechas = [], [], {}
    for campo in ('desde', 'hasta'):
        if filtros[campo]:
            try:
                if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', filtros[campo]):
                    raise ValueError()
                fecha = datetime.strptime(filtros[campo], '%Y-%m-%d')
                fechas[campo] = fecha
                # Fechas de creación inclusivas en Perú; en BD se guarda UTC.
                fecha_utc = fecha + timedelta(hours=5, days=1 if campo == 'hasta' else 0)
            except (ValueError, OverflowError):
                abort(400, description='Usa fechas válidas con formato AAAA-MM-DD.')
            condiciones.append('t.created_at >= %s' if campo == 'desde' else 't.created_at < %s')
            parametros.append(fecha_utc)
    if 'desde' in fechas and 'hasta' in fechas and fechas['desde'] > fechas['hasta']:
        abort(400, description='La fecha desde no puede ser posterior a la fecha hasta.')
    for campo, permitidos, columna in [('estado', ESTADOS, 't.estado'),
                                       ('prioridad', ('Alta', 'Media', 'Baja'), 'p.nombre')]:
        if filtros[campo]:
            if filtros[campo] not in permitidos:
                abort(400, description=f'Filtro de {campo} inválido.')
            condiciones.append(columna + '=%s')
            parametros.append(filtros[campo])
    where = ' WHERE ' + ' AND '.join(condiciones) if condiciones else ''
    return filtros, where, parametros


def resumen_analisis(where, parametros):
    # Columnas fijas del programa; no se interpolan nombres enviados por el usuario.
    grupos = [('por_estado', 't.estado', 't.estado'),
              ('por_prioridad', 'p.nombre', 'p.id,p.nombre'),
              ('por_responsable', "COALESCE(CONCAT(te.nombre,' ',te.apellido),'Sin asignar')", 'te.id,te.nombre,te.apellido'),
              ('por_categoria', 'c.nombre', 'c.id,c.nombre')]
    resumen = {}
    with obtener_conexion() as objCon:
        with objCon.cursor() as cursor:
            for nombre, etiqueta, agrupacion in grupos:
                sql = 'SELECT ' + etiqueta + ' AS etiqueta, COUNT(*) AS cantidad' + SQL_ANALISIS + where
                cursor.execute(sql + ' GROUP BY ' + agrupacion + ' ORDER BY cantidad DESC, etiqueta', parametros)
                resumen[nombre] = cursor.fetchall()
    resumen['total'] = sum(fila['cantidad'] for fila in resumen['por_estado'])
    return resumen


def filas_analisis(where, parametros, limite, desplazamiento=0):
    sql = '''SELECT t.id,t.asunto,t.descripcion,t.estado,t.created_at,t.fecha_limite,t.fecha_resolucion,
        p.nombre AS prioridad,c.nombre AS categoria,
        CONCAT(u.nombre,' ',u.apellido) AS solicitante,
        CONCAT(te.nombre,' ',te.apellido) AS tecnico''' + SQL_ANALISIS + where
    filas = consultar(sql + ' ORDER BY t.created_at DESC,t.id DESC LIMIT %s OFFSET %s',
                      list(parametros) + [limite, desplazamiento])
    momento = ahora()
    for fila in filas:
        fila['sla'] = calcular_sla(fila, momento)
    return filas


@app.route('/dashboard')
@proteger('Administrador')
def dashboard():
    # M4: dashboard junto al SLA y al semáforo.
    filtros, where, parametros = filtros_analisis()
    resumen = resumen_analisis(where, parametros)
    return render_template('dashboard_m4.html', es_reporte=False, resumen=resumen, filtros=filtros)


@app.route('/reportes')
@proteger('Administrador')
def reportes():
    # M5: reportes junto a notificaciones y API de consulta.
    filtros, where, parametros = filtros_analisis()
    total = consultar('SELECT COUNT(*) AS cantidad' + SQL_ANALISIS + where, parametros, True)['cantidad']
    filas = filas_analisis(where, parametros, 50)
    return render_template('dashboard_m4.html', es_reporte=True, tickets=filas, total=total, filtros=filtros)



def datos_reporte(tickets):
    return [[t['id'], t['asunto'], t['categoria'], t['prioridad'], t['estado'],
             t['tecnico'] or 'Sin asignar', fecha_peru(t['created_at']), t['sla']['texto']]
            for t in tickets]


def descripcion_filtros(filtros):
    return ' · '.join(f'{k.capitalize()}: {v}' for k, v in filtros.items() if v) or 'Todos los tickets'


def crear_excel(tickets, filtros):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Tickets'
    hoja.append(['Nexo TI · Reporte de tickets'])
    hoja.append([descripcion_filtros(filtros)])
    hoja.append(['Fechas en hora de Perú (UTC-5)'])
    hoja.append(['ID', 'Asunto', 'Categoría', 'Prioridad', 'Estado', 'Responsable', 'Registrado', 'SLA'])
    for fila in datos_reporte(tickets):
        hoja.append(fila)
    for fila in hoja:
        for celda in fila:
            if isinstance(celda.value, str):
                celda.data_type = 's'  # Nunca ejecutar texto del usuario como fórmula.
            celda.alignment = Alignment(vertical='top', wrap_text=True)
    for celda in hoja[4]:
        celda.fill = PatternFill('solid', fgColor='1E293B')
        celda.font = Font(color='FFFFFF', bold=True)
    for i, ancho in enumerate([8, 55, 20, 14, 18, 32, 24, 28], 1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = 'A5'
    hoja.auto_filter.ref = f'A4:H{hoja.max_row}'
    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


def crear_pdf(tickets, filtros):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    salida = BytesIO()
    estilos = getSampleStyleSheet()
    estilos['BodyText'].fontSize = 8
    estilos['BodyText'].leading = 11
    def parrafo(valor):
        return Paragraph(escape(str(valor)), estilos['BodyText'])
    contenido = [Paragraph('Nexo TI · Reporte de tickets', estilos['Title']),
                 parrafo(descripcion_filtros(filtros)),
                 parrafo(f'{len(tickets)} ticket(s) · Fechas en hora de Perú (UTC-5)'), Spacer(1, 14)]
    if tickets:
        filas = [['ID', 'Asunto', 'Categoría', 'Prioridad', 'Estado', 'Responsable', 'Registrado', 'SLA']] + datos_reporte(tickets)
        tabla = LongTable([[parrafo(v) for v in fila] for fila in filas],
                          colWidths=[32, 200, 80, 52, 70, 130, 100, 70], repeatRows=1)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DBEAFE')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
        contenido.append(tabla)
    else:
        contenido.append(parrafo('No hay tickets para los filtros seleccionados.'))
    def pie(canvas, doc):
        canvas.setFont('Helvetica', 8)
        canvas.drawRightString(landscape(A4)[0]-28, 16, f'Nexo TI | Página {doc.page}')
    SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=28, rightMargin=28,
                      topMargin=28, bottomMargin=32).build(contenido, onFirstPage=pie, onLaterPages=pie)
    salida.seek(0)
    return salida


@app.route('/reportes/exportar/<formato>')
@proteger('Administrador')
def exportar_reporte(formato):
    if formato not in ('xlsx', 'pdf'):
        abort(404)
    filtros, where, parametros = filtros_analisis()
    filas = filas_analisis(where, parametros, 1001)
    if len(filas) > 1000:
        abort(400, description='El reporte supera los 1000 tickets. Reduce el período o aplica filtros.')
    archivo = crear_excel(filas, filtros) if formato == 'xlsx' else crear_pdf(filas, filtros)
    mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if formato == 'xlsx' else 'application/pdf'
    return send_file(archivo, mimetype=mime, as_attachment=True,
                     download_name=f'tickets_{ahora():%Y%m%d}.{formato}', max_age=0)


@app.route('/api/tickets')
@proteger('Administrador')
def api_tickets():
    # M5: API de consulta; conserva la protección y paginación del primer avance.
    filtros, where, parametros = filtros_analisis()
    try:
        pagina = int(request.args.get('pagina', '1'))
        limite = int(request.args.get('limite', '20'))
        if not 1 <= pagina <= 1000000 or not 1 <= limite <= 100:
            raise ValueError()
    except ValueError:
        abort(400, description='pagina debe estar entre 1 y 1000000 y limite entre 1 y 100.')
    total = consultar('SELECT COUNT(*) AS cantidad' + SQL_ANALISIS + where, parametros, True)['cantidad']
    filas = filas_analisis(where, parametros, limite, (pagina - 1) * limite)
    return jsonify(datos=[ticket_para_json(t) for t in filas], total=total,
                   pagina=pagina, limite=limite, filtros=filtros)


def ticket_para_json(ticket):
    # Copia explícita, sin correos, contraseñas, rutas de fotos ni datos de sesión.
    resultado = dict(ticket)
    for campo in ('created_at', 'fecha_limite', 'fecha_resolucion'):
        fecha = resultado[campo]
        resultado[campo] = fecha.isoformat(timespec='seconds') + 'Z' if fecha else None
    return resultado


@app.route('/api/tickets/<int:ticket_id>')
@proteger('Administrador')
def api_ticket_detalle(ticket_id):
    filas = filas_analisis(' WHERE t.id=%s', [ticket_id], 1)
    if not filas:
        abort(404)
    return jsonify(datos=ticket_para_json(filas[0]))


@app.route('/api/resumen')
@proteger('Administrador')
def api_resumen():
    filtros, where, parametros = filtros_analisis()
    return jsonify(datos=resumen_analisis(where, parametros), filtros=filtros)


@app.after_request
def privacidad_analisis(respuesta):
    if request.path.startswith(('/api/', '/reportes', '/dashboard')):
        respuesta.headers['Cache-Control'] = 'private, no-store'
    return respuesta


@app.errorhandler(400)
@app.errorhandler(401)
@app.errorhandler(403)
@app.errorhandler(404)
@app.errorhandler(405)
@app.errorhandler(413)
def error_http(error):
    mensajes = {400: error.description, 403: 'No tienes permiso para acceder a este recurso.',
                401: 'Inicia sesión con una cuenta activa.', 405: 'Método no permitido.',
                404: 'No encontramos el recurso solicitado.', 413: 'Los adjuntos superan el límite total de 16 MB.'}
    if request.path.startswith('/api/'):
        return jsonify(error=dict(codigo=error.code, mensaje=mensajes[error.code])), error.code
    return render_template('error_comun.html', codigo=error.code, mensaje=mensajes[error.code]), error.code


@app.errorhandler(pymysql.MySQLError)
def error_bd(error):
    app.logger.error('Error de base de datos: %s', error)
    # Evitar que el procesador de contexto repita la consulta cuando la BD falla.
    g.usuario = None
    if request.path.startswith('/api/'):
        return jsonify(error=dict(codigo=503, mensaje='Base de datos no disponible.')), 503
    return render_template('error_comun.html', codigo=503,
                           mensaje='No se pudo acceder a la base de datos. Revisa que MySQL esté iniciado y que se hayan importado los scripts.'), 503


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000, debug=False)
